"""DOCX text and table extraction using python-docx."""
from __future__ import annotations

from pathlib import Path

import docx

from src.models.schemas import ExtractionResult, PageText


class DocxExtractor:
    """Extracts text and tables from Word (.docx) files.

    Note: DOCX files don't have a native concept of pages, so all content
    is returned as page 1. Table data is extracted with structure preserved.
    """

    def __init__(self, path: str | Path):
        self.path = Path(path)

    def extract(self) -> ExtractionResult:
        doc = docx.Document(str(self.path))

        paragraphs: list[str] = []
        tables: list[list[list[str]]] = []

        # --- Paragraphs (preserving heading hierarchy) ---
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            if para.style.name.startswith("Heading"):
                paragraphs.append(f"\n{'#' * self._heading_level(para.style.name)} {text}")
            else:
                paragraphs.append(text)

        # --- Tables ---
        for table in doc.tables:
            rows: list[list[str]] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(cells)
            if rows:
                tables.append(rows)

        raw_text = "\n".join(paragraphs)
        page = PageText(page_number=1, text=raw_text, tables=tables)

        return ExtractionResult(
            source_file=str(self.path),
            total_pages=1,
            pages=[page],
            raw_text=raw_text,
        )

    @staticmethod
    def _heading_level(style_name: str) -> int:
        """Extract numeric level from style name e.g. 'Heading 2' → 2."""
        parts = style_name.split()
        if len(parts) >= 2 and parts[-1].isdigit():
            return int(parts[-1])
        return 1
